from crewai_tools import BaseTool
from docx import Document
from fpdf import FPDF
import jinja2
from jinja2 import Template

# Tool for writing to DOCX files
class DOCXWriteTool(BaseTool):
    name = "DOCXWriteTool"
    description = "Writes provided text to a DOCX file."

    def _run(self, file_path, text):
        doc = Document()
        doc.add_paragraph(text)
        doc.save(file_path)
        return f"Text written successfully to {file_path}."

# Tool for writing to PDF files
class PDFWriteTool(BaseTool):
    name = "PDFWriteTool"
    description = "Creates a PDF file and writes text to it."

    def _run(self, file_path, text):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size = 12)
        pdf.cell(200, 10, txt = text, ln = True, align = 'C')
        pdf.output(file_path)
        return f"Text written successfully to {file_path}."

# Tool for writing to HTML files
class HTMLWriteTool(BaseTool):
    name = "HTMLWriteTool"
    description = "Generates an HTML file from provided text."

    def _run(self, file_path, text):
        template = Template("<html><head><title>Output</title></head><body><p>{{ text }}</p></body></html>")
        html_content = template.render(text=text)
        with open(file_path, "w") as file:
            file.write(html_content)
        return f"HTML content written successfully to {file_path}."

# Example
docx_tool = DOCXWriteTool()
pdf_tool = PDFWriteTool()
html_tool = HTMLWriteTool()

docx_tool.run('output.docx', 'Hello, this is some text for a DOCX file.')
pdf_tool.run('output.pdf', 'Hello, this is some text for a PDF file.')
html_tool.run('output.html', 'Hello, this is some text for an HTML file.')